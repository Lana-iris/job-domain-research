#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
coverage_tool.py — 平台覆盖台账维护工具（通用）

只做一件事：把「今天查了哪些平台、哪些检索位、查出多少条、入库多少条」
写进一张可审计的台账。不联网、不搜索、不写 JD 库。

用法：
    python3 coverage_tool.py --workbook <路径> append rows.json [--dry-run] [--allow-duplicate]
    python3 coverage_tool.py --workbook <路径> update --date 2026-10-21 --platform 猎聘 \
        --qp QP1 --city 深圳 --inbound 3 [--dry-run]
    python3 coverage_tool.py --workbook <路径> audit --date 2026-10-21 [--plan plan.json]
    python3 coverage_tool.py describe [--format markdown|json]

    --workbook 默认当前目录下的「platform-coverage.xlsx」。

设计要点
- 台账字段规格以本文件 COVERAGE_SPEC 为唯一真源，describe 直接由它生成。
- 列一律按表头名定位，不按列号硬编码。
- 任何写操作前自动备份到工作簿同目录下的 `备份/`。
- 校验失败整批拒绝，不写半批数据。
- audit 只读，不产生备份。
"""

from __future__ import annotations

import argparse
import datetime as dt
import json
import re
import shutil
import sys
from pathlib import Path
from typing import Any

try:
    import openpyxl
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover
    sys.exit("需要 openpyxl：pip install openpyxl")


# ---------------------------------------------------------------- 字段规格（唯一真源）

COVERAGE_SPEC: list[dict[str, Any]] = [
    {
        "name": "检索日期",
        "meaning": "检索执行日期",
        "required": True,
        "type": "YYYY-MM-DD",
        "write_rule": "必须是真实日期，统一写成 YYYY-MM-DD",
        "example": "2026-10-21",
    },
    {
        "name": "平台",
        "meaning": "实际检索的平台",
        "required": True,
        "type": "文本",
        "write_rule": "写平台本名，如猎聘、智联招聘、BOSS直聘",
        "example": "猎聘",
    },
    {
        "name": "检索位/查询词",
        "meaning": "检索位 ID + 实际查询词",
        "required": True,
        "type": "QP+数字｜查询词",
        "write_rule": "必须以 QP 编号开头，用全角｜或半角|分隔查询词",
        "example": "QP1｜AI产品经理 人力资源 深圳",
    },
    {
        "name": "城市",
        "meaning": "本条查询覆盖的城市或地域",
        "required": True,
        "type": "文本",
        "write_rule": "写实际检索地域；全国检索写「全国」",
        "example": "深圳",
    },
    {
        "name": "检索状态",
        "meaning": "本次检索的执行结果",
        "required": True,
        "type": "已检索 / 受限 / 未执行",
        "write_rule": "三选一；受限和未执行必须在说明列写清原因",
        "example": "已检索",
    },
    {
        "name": "相关候选数",
        "meaning": "本次查询实际识别出的相关职位线索数",
        "required": "已检索/受限必填；未执行可空",
        "type": "整数 ≥ 0",
        "write_rule": "只数真正相关的线索，不把搜索结果总数当候选数",
        "example": "8",
    },
    {
        "name": "入库数",
        "meaning": "本次扫描后实际新增进活跃 JD 库的岗位数",
        "required": False,
        "type": "整数 ≥ 0",
        "write_rule": "扫描阶段可留空或写 0；O2b 入库后回填；未执行时只能留空或 0；不得大于相关候选数",
        "example": "3",
    },
    {
        "name": "访问限制/筛选说明",
        "meaning": "登录墙、详情不可访问、地域不符、冲突、过期、未执行原因等",
        "required": "未执行/受限必填；已检索且候选数为 0 必填",
        "type": "文本",
        "write_rule": "写可复核的具体原因；零候选优先写「未发现可核验结果」再补原因",
        "example": "未发现可核验结果：搜索结果均为销售岗",
    },
    {
        "name": "证据链接",
        "meaning": "一条最能代表本行结论的搜索页或职位详情链接",
        "required": False,
        "type": "http(s) URL",
        "write_rule": "可选；填就必须以 http:// 或 https:// 开头",
        "example": "https://www.liepin.com/zhaopin/?key=AI产品经理",
    },
]

FIELD_NAMES = [f["name"] for f in COVERAGE_SPEC]
KEY_FIELDS = ("检索日期", "平台", "检索位/查询词", "城市")
STATUSES = ("已检索", "受限", "未执行")
QP_PATTERN = re.compile(r"^QP\d+\s*[｜|].+")
QP_ID_PATTERN = re.compile(r"^(QP\d+)")
PLAN_QP_ID_PATTERN = re.compile(r"^QP\d+$")
DATE_PATTERN = re.compile(r"^\d{4}-\d{2}-\d{2}$")

DEFAULT_WORKBOOK = "platform-coverage.xlsx"
DEFAULT_SHEET = "平台覆盖台账"


# ---------------------------------------------------------------- 基础工具


def resolve_workbook(raw: str | None) -> Path:
    return Path(raw).expanduser() if raw else Path.cwd() / DEFAULT_WORKBOOK


def is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and value.strip() == "")


def normalize_date(value: Any) -> str:
    """把工作表里的单元格值统一成 YYYY-MM-DD 字符串。"""
    if isinstance(value, dt.datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, dt.date):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, str):
        return value.strip()
    if value is None:
        return ""
    return str(value).strip()


def normalize_cell(field: str, value: Any) -> Any:
    if field == "检索日期":
        return normalize_date(value)
    if isinstance(value, str):
        return value.strip()
    return value


def qp_id(query: Any) -> str:
    match = QP_ID_PATTERN.match(str(query or "").strip())
    return match.group(1) if match else ""


def row_key(row: dict[str, Any]) -> tuple[str, str, str, str]:
    return tuple(str(row.get(name) or "").strip() for name in KEY_FIELDS)  # type: ignore[return-value]


def parse_int(value: Any, field: str, errors: list[str]) -> int | None:
    """严格解析非负整数：只接受 JSON 整数，拒绝 bool / float / 文本。"""
    if isinstance(value, bool):
        errors.append(f"{field} 必须是整数，不能是布尔值")
        return None
    if isinstance(value, int):
        if value < 0:
            errors.append(f"{field} 不能为负数（当前 {value}）")
            return None
        return value
    errors.append(f"{field} 必须是 JSON 整数，不能是文本、小数或空值（当前 {value!r}）")
    return None


def load_sheet(path: Path):
    if not path.exists():
        raise SystemExit(f"找不到工作簿：{path}")
    wb = openpyxl.load_workbook(path)
    if DEFAULT_SHEET in wb.sheetnames:
        ws = wb[DEFAULT_SHEET]
    else:
        ws = wb.active
        print(
            f"[警告] 工作簿没有「{DEFAULT_SHEET}」sheet，改用「{ws.title}」",
            file=sys.stderr,
        )
    return wb, ws


def header_map(ws) -> dict[str, int]:
    """表头名 -> 列号（1-based）；重名直接报错。"""
    mapping: dict[str, int] = {}
    duplicates: list[str] = []
    for cell in ws[1]:
        name = normalize_cell("", cell.value)
        if not isinstance(name, str):
            name = str(name)
        name = name.strip()
        if not name:
            continue
        if name in mapping:
            duplicates.append(name)
        else:
            mapping[name] = cell.column
    if duplicates:
        raise SystemExit(f"台账表头存在重复列名：{'、'.join(sorted(set(duplicates)))}")
    return mapping


def check_headers(ws) -> dict[str, int]:
    mapping = header_map(ws)
    missing = [name for name in FIELD_NAMES if name not in mapping]
    if missing:
        raise SystemExit(f"台账表头缺少必需列：{'、'.join(missing)}")
    extra = [name for name in mapping if name not in FIELD_NAMES]
    if extra:
        print(f"[警告] 台账存在额外列（不会写入）：{'、'.join(extra)}", file=sys.stderr)
    return mapping


def backup_workbook(path: Path) -> Path:
    backup_dir = path.parent / "备份"
    backup_dir.mkdir(parents=True, exist_ok=True)
    stamp = dt.datetime.now().strftime("%Y%m%d-%H%M%S")
    target = backup_dir / f"{path.stem}.bak-{stamp}{path.suffix}"
    counter = 1
    while target.exists():
        target = backup_dir / f"{path.stem}.bak-{stamp}-{counter}{path.suffix}"
        counter += 1
    shutil.copy2(path, target)
    print(f"  [备份] {target}")
    return target


def read_rows(ws, mapping: dict[str, int]) -> list[tuple[int, dict[str, Any]]]:
    """读取数据行（跳过全空行），按字段名归一化。"""
    rows: list[tuple[int, dict[str, Any]]] = []
    for row_idx in range(2, ws.max_row + 1):
        values = {
            name: normalize_cell(name, ws.cell(row=row_idx, column=col).value)
            for name, col in mapping.items()
            if name in FIELD_NAMES
        }
        if all(is_blank(values.get(name)) for name in FIELD_NAMES):
            continue
        rows.append((row_idx, values))
    return rows


# ---------------------------------------------------------------- 校验


def validate_item(raw: Any, index: int) -> tuple[dict[str, Any] | None, list[str], list[str]]:
    """返回 (规范化行, 错误, 警告)。错误非空表示该行不可写。"""
    errors: list[str] = []
    warnings: list[str] = []
    label = f"第 {index} 条"

    if not isinstance(raw, dict):
        return None, [f"{label}：必须是 JSON 对象"], warnings

    unknown = [key for key in raw if key not in FIELD_NAMES]
    if unknown:
        errors.append(f"{label}：未知字段 {'、'.join(map(str, unknown))}（只接受台账 9 列字段名）")

    row: dict[str, Any] = {name: raw.get(name) for name in FIELD_NAMES}
    for name in FIELD_NAMES:
        if isinstance(row[name], str):
            row[name] = row[name].strip()

    # 必填文本
    for name in ("平台", "城市"):
        value = row[name]
        if is_blank(value):
            errors.append(f"{label}：{name} 必填")
        elif not isinstance(value, str):
            errors.append(f"{label}：{name} 必须是文本，不能是数字或其他类型（当前 {value!r}）")

    # 检索日期
    if is_blank(row["检索日期"]):
        errors.append(f"{label}：检索日期 必填")
    else:
        text = str(row["检索日期"])
        if not DATE_PATTERN.match(text):
            errors.append(f"{label}：检索日期 必须是 YYYY-MM-DD（当前 {text!r}）")
        else:
            try:
                parsed = dt.datetime.strptime(text, "%Y-%m-%d").date()
            except ValueError:
                errors.append(f"{label}：检索日期 不是真实日期（当前 {text!r}）")
            else:
                row["检索日期"] = parsed.strftime("%Y-%m-%d")

    # 检索位/查询词
    query = row["检索位/查询词"]
    if is_blank(query):
        errors.append(f"{label}：检索位/查询词 必填")
    elif not QP_PATTERN.match(str(query)):
        errors.append(f"{label}：检索位/查询词 必须以 QP 编号开头并用｜或|分隔查询词（当前 {query!r}）")

    # 检索状态
    status = row["检索状态"]
    if is_blank(status):
        errors.append(f"{label}：检索状态 必填")
    elif status not in STATUSES:
        errors.append(f"{label}：检索状态 只能是 {' / '.join(STATUSES)}（当前 {status!r}）")

    # 数字列
    candidates = row["相关候选数"]
    inbound = row["入库数"]
    candidates_int: int | None = None
    inbound_int: int | None = None

    if is_blank(candidates):
        if status in ("已检索", "受限"):
            errors.append(f"{label}：检索状态为 {status} 时，相关候选数 必填")
    else:
        candidates_int = parse_int(candidates, "相关候选数", errors)

    if not is_blank(inbound):
        inbound_int = parse_int(inbound, "入库数", errors)
    if candidates_int is not None and inbound_int is not None and inbound_int > candidates_int:
        errors.append(f"{label}：入库数（{inbound_int}）不能大于相关候选数（{candidates_int}）")

    # 说明列的条件必填
    reason = row["访问限制/筛选说明"]
    if status == "未执行":
        if inbound_int not in (None, 0):
            errors.append(f"{label}：未执行时 入库数 只能留空或写 0")
        if is_blank(reason):
            errors.append(f"{label}：未执行时 访问限制/筛选说明 必填（写清未执行原因）")
    elif status == "受限":
        if is_blank(reason):
            errors.append(f"{label}：受限时 访问限制/筛选说明 必填（写清限制类型）")
    elif status == "已检索" and candidates_int == 0:
        if is_blank(reason):
            errors.append(f"{label}：已检索且相关候选数为 0 时，访问限制/筛选说明 必填")
        elif "未发现可核验结果" not in str(reason):
            warnings.append(f"{label}：零候选说明未包含「未发现可核验结果」，请确认已写清原因")

    # 证据链接
    evidence = row["证据链接"]
    if not is_blank(evidence):
        text = str(evidence)
        if not (text.startswith("http://") or text.startswith("https://")):
            errors.append(f"{label}：证据链接 必须以 http:// 或 https:// 开头（当前 {text!r}）")

    # 落盘值：空字符串视为空单元格
    clean = {
        name: (None if is_blank(row[name]) else row[name])
        for name in FIELD_NAMES
    }
    if candidates_int is not None:
        clean["相关候选数"] = candidates_int
    if inbound_int is not None:
        clean["入库数"] = inbound_int

    return (None if errors else clean), errors, warnings


def load_json_items(path: Path) -> list[Any]:
    if not path.exists():
        raise SystemExit(f"找不到 JSON 文件：{path}")
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as exc:
        raise SystemExit(f"JSON 解析失败：{exc}") from exc
    if isinstance(data, dict):
        return [data]
    if isinstance(data, list):
        return data
    raise SystemExit("JSON 顶层必须是对象或对象数组")


# ---------------------------------------------------------------- append


def cmd_append(args: argparse.Namespace) -> int:
    path = resolve_workbook(args.workbook)
    wb, ws = load_sheet(path)
    mapping = check_headers(ws)
    items = load_json_items(Path(args.input).expanduser())

    if not items:
        print("已追加 0 条；跳过/拒绝 0 条（输入为空数组，未写入）")
        return 0

    existing = {row_key(values) for _, values in read_rows(ws, mapping)}
    seen: dict[tuple[str, str, str, str], int] = {}

    accepted: list[dict[str, Any]] = []
    all_errors: list[str] = []
    all_warnings: list[str] = []

    for index, raw in enumerate(items, start=1):
        row, errors, warnings = validate_item(raw, index)
        all_warnings.extend(warnings)
        if row is not None:
            key = row_key(row)
            where = seen.get(key)
            if key in existing and not args.allow_duplicate:
                errors.append(
                    f"第 {index} 条：重复记录（{key[0]} / {key[1]} / {key[2]} / {key[3]}）已存在"
                )
            elif where is not None:
                errors.append(f"第 {index} 条：与本批第 {where} 条重复")
            else:
                seen[key] = index
        if errors:
            all_errors.extend(errors)
        elif row is not None:
            accepted.append(row)

    for warning in all_warnings:
        print(f"[警告] {warning}", file=sys.stderr)

    if all_errors:
        for error in all_errors:
            print(f"[拒绝] {error}", file=sys.stderr)
        print(f"已追加 0 条；跳过/拒绝 {len(items)} 条（整批未写入）", file=sys.stderr)
        return 1

    if args.dry_run:
        for row in accepted:
            print(f"  [预演] {row['检索日期']} / {row['平台']} / {row['检索位/查询词']} / {row['城市']}")
        print(f"已追加 0 条；跳过/拒绝 0 条（--dry-run，未写入）")
        return 0

    backup_workbook(path)
    for row in accepted:
        row_idx = ws.max_row + 1
        for name, col in mapping.items():
            if name in FIELD_NAMES:
                ws.cell(row=row_idx, column=col, value=row[name])
    wb.save(path)
    print(f"已追加 {len(accepted)} 条；跳过/拒绝 0 条")
    return 0


# ---------------------------------------------------------------- update


def parse_cli_int(raw: Any, flag: str) -> int:
    """CLI 参数本身都是字符串；只接受纯数字，转成 int 后再按整数规则校验。"""
    text = str(raw).strip()
    if not re.fullmatch(r"\d+", text):
        raise SystemExit(f"{flag} 必须是非负整数（当前 {raw!r}）")
    return int(text)


def cmd_update(args: argparse.Namespace) -> int:
    path = resolve_workbook(args.workbook)
    wb, ws = load_sheet(path)
    mapping = check_headers(ws)

    date_text = normalize_date(args.date)
    try:
        dt.datetime.strptime(date_text, "%Y-%m-%d")
    except ValueError:
        raise SystemExit(f"--date 必须是真实日期 YYYY-MM-DD（当前 {args.date!r}）")

    inbound = parse_cli_int(args.inbound, "--inbound")

    matches: list[tuple[int, dict[str, Any]]] = []
    for row_idx, values in read_rows(ws, mapping):
        if values.get("检索日期") != date_text:
            continue
        if str(values.get("平台") or "") != args.platform:
            continue
        if str(values.get("城市") or "") != args.city:
            continue
        query = str(values.get("检索位/查询词") or "")
        if query != args.qp and qp_id(query) != args.qp:
            continue
        matches.append((row_idx, values))

    if not matches:
        raise SystemExit(
            f"找不到匹配行：{date_text} / {args.platform} / {args.qp} / {args.city}"
        )
    if len(matches) > 1:
        raise SystemExit(f"匹配到 {len(matches)} 行，无法确定更新目标，请用完整查询词区分")

    row_idx, values = matches[0]
    if values.get("检索状态") == "未执行" and inbound > 0:
        raise SystemExit("未执行的行 入库数 只能留空或写 0")

    candidates = values.get("相关候选数")
    if not is_blank(candidates):
        errors: list[str] = []
        candidates_int = parse_int(candidates, "相关候选数", errors)
        if candidates_int is None:
            raise SystemExit(f"该行 相关候选数 无效：{errors[0]}")
        if inbound > candidates_int:
            raise SystemExit(f"入库数（{inbound}）不能大于相关候选数（{candidates_int}）")

    old = values.get("入库数")
    old_display = "空" if is_blank(old) else str(old)
    print(f"  行 {row_idx}：{date_text} / {args.platform} / {values.get('检索位/查询词')} / {args.city}")
    print(f"  入库数：{old_display} → {inbound}")

    if args.dry_run:
        print("  （--dry-run，未写入）")
        return 0

    backup_workbook(path)
    ws.cell(row=row_idx, column=mapping["入库数"], value=inbound)
    wb.save(path)
    return 0


# ---------------------------------------------------------------- audit


def load_plan(path: Path) -> dict[str, Any]:
    """载入并校验覆盖计划；所有错误都转成干净的中文报错。"""
    if not path.exists():
        raise SystemExit(f"找不到 plan 文件：{path}")
    try:
        text = path.read_text(encoding="utf-8")
    except OSError as exc:
        raise SystemExit(f"计划文件读取失败：{exc}") from exc
    try:
        data = json.loads(text)
    except json.JSONDecodeError as exc:
        raise SystemExit(f"计划文件 JSON 解析失败：{exc}") from exc

    if not isinstance(data, dict):
        raise SystemExit("计划文件顶层必须是 JSON 对象")

    if "platforms" not in data:
        print("[警告] 计划文件缺少 platforms，按空列表处理", file=sys.stderr)
        platforms: list[str] = []
    else:
        raw_platforms = data["platforms"]
        if not isinstance(raw_platforms, list):
            raise SystemExit("计划文件 platforms 必须是字符串数组")
        platforms = []
        for index, value in enumerate(raw_platforms, start=1):
            if not isinstance(value, str) or not value.strip():
                raise SystemExit(
                    f"计划文件 platforms 第 {index} 项必须是非空字符串（当前 {value!r}）"
                )
            platforms.append(value.strip())

    if "qps" not in data:
        print("[警告] 计划文件缺少 qps，按空列表处理", file=sys.stderr)
        qps: list[dict[str, Any]] = []
    else:
        raw_qps = data["qps"]
        if not isinstance(raw_qps, list):
            raise SystemExit("计划文件 qps 必须是对象数组")
        qps = []
        seen: set[str] = set()
        for index, item in enumerate(raw_qps, start=1):
            if not isinstance(item, dict):
                raise SystemExit(f"计划文件 qps 第 {index} 项必须是 JSON 对象")
            qp = item.get("id")
            if not isinstance(qp, str) or not PLAN_QP_ID_PATTERN.match(qp.strip()):
                raise SystemExit(
                    f"计划文件 qps 第 {index} 项的 id 必须匹配 QP+数字（当前 {qp!r}）"
                )
            qp = qp.strip()
            if qp in seen:
                raise SystemExit(f"计划文件 qps 存在重复 id：{qp}")
            seen.add(qp)
            minimum = item.get("min_queries")
            if isinstance(minimum, bool) or not isinstance(minimum, int) or minimum < 0:
                raise SystemExit(
                    f"计划文件 qps {qp} 的 min_queries 必须是非负整数（当前 {minimum!r}）"
                )
            qps.append({"id": qp, "min_queries": minimum})

    return {"platforms": platforms, "qps": qps}


def dedupe_rows(
    rows: list[tuple[int, dict[str, Any]]],
) -> tuple[list[tuple[int, dict[str, Any]]], list[tuple[int, dict[str, Any]]]]:
    """按 key 去重；保留首次出现，返回 (有效行, 被忽略的重复行)。"""
    kept: list[tuple[int, dict[str, Any]]] = []
    duplicates: list[tuple[int, dict[str, Any]]] = []
    seen: set[tuple[str, str, str, str]] = set()
    for row_idx, values in rows:
        key = row_key(values)
        if key in seen:
            duplicates.append((row_idx, values))
        else:
            seen.add(key)
            kept.append((row_idx, values))
    return kept, duplicates


def audit_int(
    value: Any,
    row_idx: int,
    field: str,
    col: int,
    warnings: list[str],
) -> int | None:
    """审计用整数读取：无效值不静默跳过，记一条带行/单元格的警告。"""
    if is_blank(value):
        return None
    if isinstance(value, bool) or not isinstance(value, int) or value < 0:
        coordinate = f"{get_column_letter(col)}{row_idx}"
        warnings.append(
            f"第 {row_idx} 行 {field}（{coordinate}）不是非负整数：{value!r}，未计入合计"
        )
        return None
    return value


def is_zero_candidate(value: Any) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool) and value == 0


def cmd_audit(args: argparse.Namespace) -> int:
    path = resolve_workbook(args.workbook)
    wb, ws = load_sheet(path)
    mapping = check_headers(ws)

    date_text = normalize_date(args.date)
    try:
        dt.datetime.strptime(date_text, "%Y-%m-%d")
    except ValueError:
        raise SystemExit(f"--date 必须是真实日期 YYYY-MM-DD（当前 {args.date!r}）")

    plan = load_plan(Path(args.plan).expanduser()) if args.plan else None

    raw_rows = [
        (row_idx, values)
        for row_idx, values in read_rows(ws, mapping)
        if values.get("检索日期") == date_text
    ]
    rows, duplicates = dedupe_rows(raw_rows)

    status_count: dict[str, int] = {}
    platforms: dict[str, dict[str, Any]] = {}
    per_qp_queries: dict[str, set[str]] = {}
    numeric_warnings: list[str] = []

    for row_idx, values in rows:
        status = str(values.get("检索状态") or "未填")
        status_count[status] = status_count.get(status, 0) + 1

        platform = str(values.get("平台") or "未填")
        entry = platforms.setdefault(
            platform,
            {"rows": 0, "statuses": {}, "candidates": 0, "inbound": 0},
        )
        entry["rows"] += 1
        entry["statuses"][status] = entry["statuses"].get(status, 0) + 1

        candidates = audit_int(
            values.get("相关候选数"), row_idx, "相关候选数", mapping["相关候选数"], numeric_warnings
        )
        if candidates is not None:
            entry["candidates"] += candidates
        inbound = audit_int(
            values.get("入库数"), row_idx, "入库数", mapping["入库数"], numeric_warnings
        )
        if inbound is not None:
            entry["inbound"] += inbound

        qp = qp_id(values.get("检索位/查询词")) or "未标注"
        per_qp_queries.setdefault(qp, set()).add(
            str(values.get("检索位/查询词") or "")
        )

    per_qp = {qp: len(queries) for qp, queries in per_qp_queries.items()}

    print(f"覆盖台账审计 · {date_text}")
    print(f"工作簿：{path}")
    print("=" * 48)

    if not rows:
        print("该日期没有记录。")
    else:
        status_line = "，".join(
            f"{key} {value}" for key, value in sorted(status_count.items())
        )
        print(f"总行数：{len(rows)}    状态：{status_line}")

        print("\n按平台：")
        for platform in sorted(platforms):
            entry = platforms[platform]
            statuses = "，".join(
                f"{key} {value}" for key, value in sorted(entry["statuses"].items())
            )
            print(
                f"  {platform}：{entry['rows']} 行（{statuses}）"
                f"  候选 {entry['candidates']}  入库 {entry['inbound']}"
            )

        print("\n按检索位：")
        for qp in sorted(per_qp, key=lambda item: (not item.startswith("QP"), item)):
            count = per_qp[qp]
            if plan is not None and plan["qps"]:
                minimum = next(
                    (item["min_queries"] for item in plan["qps"] if item["id"] == qp),
                    0,
                )
                verdict = "通过" if count >= minimum else "不足"
                print(f"  {qp}：{count} 个不同查询，最低要求 {minimum}，{verdict}")
            else:
                print(f"  {qp}：{count} 个不同查询")

        restricted = [values for _, values in rows if values.get("检索状态") == "受限"]
        print("\n受限行：")
        if restricted:
            for values in restricted:
                print(
                    f"  {values.get('平台')} / {values.get('检索位/查询词')} / {values.get('城市')}"
                    f" — {values.get('访问限制/筛选说明') or '未写说明'}"
                )
        else:
            print("  无")

        zero = [
            values for _, values in rows if is_zero_candidate(values.get("相关候选数"))
        ]
        print("\n零候选行：")
        if zero:
            for values in zero:
                print(
                    f"  {values.get('平台')} / {values.get('检索位/查询词')} / {values.get('城市')}"
                    f" — {values.get('访问限制/筛选说明') or '未写说明'}"
                )
        else:
            print("  无")

        pending = [
            values
            for _, values in rows
            if values.get("检索状态") == "已检索" and is_blank(values.get("入库数"))
        ]
        print(f"\n入库数待回填：{len(pending)} 行（仅统计已检索）")
        for values in pending:
            print(f"  {values.get('平台')} / {values.get('检索位/查询词')} / {values.get('城市')}")

    if duplicates:
        print("\n重复行：")
        for row_idx, values in duplicates:
            print(
                f"  第 {row_idx} 行 {values.get('平台')} / {values.get('检索位/查询词')}"
                f" / {values.get('城市')} — 与首条同键，已忽略"
            )
        print(f"  已忽略 {len(duplicates)} 行；聚合只保留首次出现")

    for warning in numeric_warnings:
        print(f"[警告] {warning}", file=sys.stderr)

    if plan is not None:
        print("\n对照 plan：")
        missing_platforms = [
            name for name in dict.fromkeys(plan["platforms"]) if name not in platforms
        ]
        if missing_platforms:
            print(f"  缺失平台：{'、'.join(missing_platforms)}")
        else:
            print("  平台覆盖：齐全")

        shortfalls = [
            f"{item['id']}（要求 {item['min_queries']} 个不同查询，实际 {per_qp.get(item['id'], 0)} 个）"
            for item in plan["qps"]
            if per_qp.get(item["id"], 0) < item["min_queries"]
        ]
        if shortfalls:
            print(f"  检索位不足：{'、'.join(shortfalls)}")
        else:
            print("  检索位覆盖：达标")
    elif rows:
        print("\n[警告] 未提供 --plan，只能看已记录内容，无法判断是否有平台或检索位漏检。")

    return 0


# ---------------------------------------------------------------- describe


def cmd_describe(args: argparse.Namespace) -> int:
    if args.format == "json":
        print(json.dumps(COVERAGE_SPEC, ensure_ascii=False, indent=2))
        return 0

    def cell(text: Any) -> str:
        return str(text).replace("|", "\\|").replace("\n", " ")

    def required_text(value: Any) -> str:
        if value is True:
            return "是"
        if value is False:
            return "否"
        return cell(value)

    print("| 字段 | 含义 | 必填 | 类型/取值 | 写入规则 | 示例 |")
    print("| --- | --- | --- | --- | --- | --- |")
    for field in COVERAGE_SPEC:
        print(
            f"| {cell(field['name'])} | {cell(field['meaning'])} | "
            f"{required_text(field['required'])} | {cell(field['type'])} | "
            f"{cell(field['write_rule'])} | {cell(field['example'])} |"
        )
    return 0


# ---------------------------------------------------------------- CLI


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        prog="coverage_tool.py",
        description="平台覆盖台账维护工具：追加检索记录、回填入库数、按日审计。",
    )
    parser.add_argument(
        "--workbook",
        help=f"台账路径（默认当前目录下的 {DEFAULT_WORKBOOK}）",
    )
    sub = parser.add_subparsers(dest="command", required=True, metavar="<command>")

    append = sub.add_parser("append", help="追加一条或多条检索记录（JSON）")
    append.add_argument("input", help="JSON 文件：对象数组或单个对象")
    append.add_argument("--dry-run", action="store_true", help="只校验和预览，不写入")
    append.add_argument(
        "--allow-duplicate",
        action="store_true",
        help="允许写入与台账已有记录同键的行",
    )
    append.set_defaults(func=cmd_append)

    update = sub.add_parser("update", help="回填某一行的入库数")
    update.add_argument("--date", required=True, help="检索日期 YYYY-MM-DD")
    update.add_argument("--platform", required=True, help="平台名")
    update.add_argument("--qp", required=True, help="检索位 ID（如 QP1）或完整查询词")
    update.add_argument("--city", required=True, help="城市")
    update.add_argument("--inbound", required=True, help="入库数（非负整数）")
    update.add_argument("--dry-run", action="store_true", help="只预览旧值 → 新值，不写入")
    update.set_defaults(func=cmd_update)

    audit = sub.add_parser("audit", help="按日期审计覆盖情况（只读）")
    audit.add_argument("--date", required=True, help="检索日期 YYYY-MM-DD")
    audit.add_argument("--plan", help="覆盖计划 JSON，用于判断平台/检索位是否漏检")
    audit.set_defaults(func=cmd_audit)

    describe = sub.add_parser("describe", help="打印字段规格")
    describe.add_argument(
        "--format",
        choices=("markdown", "json"),
        default="markdown",
        help="输出格式（默认 markdown）",
    )
    describe.set_defaults(func=cmd_describe)

    return parser


def main(argv: list[str] | None = None) -> int:
    parser = build_parser()
    args = parser.parse_args(argv)
    return args.func(args)


if __name__ == "__main__":
    sys.exit(main())
